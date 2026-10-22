#!/usr/bin/env python3
"""Будує Excel-книгу лаб. роботи №7 (варіант 9): дані, таблиця 7.4 з формулами та 3 графіки.

Використання (з кореня репозиторію):
    .venv/bin/pip install openpyxl
    .venv/bin/python analysis/make_excel.py [results/flight_segment.csv] [results/lab7_variant9.xlsx]

Вхід: flight_segment.csv, який створює analyze.py (рядки контрольного польоту від arm до disarm).
Аркуші: «Дані», «Таблиця 7.4», «Графіки».
"""
import os, sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

H, V = 35.0, 4.0
TARGET = (-35.3624857, 149.1642559)
BASE = os.path.dirname(os.path.abspath(__file__))
src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(BASE, "..", "results", "flight_segment.csv")
out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(BASE, "..", "results", "lab7_variant9.xlsx")

df = pd.read_csv(src, low_memory=False)
def col(*names):
    low = {c.lower(): c for c in df.columns}
    for n in names:
        if n.lower() in low: return low[n.lower()]
    return None
c_ts, c_alt = col("Timestamp"), col("altitudeRelative")
c_air, c_gs = col("airSpeed"), col("groundSpeed")
c_roll, c_pitch = col("roll"), col("pitch")
c_lat, c_lon = col("gps.lat", "latitude"), col("gps.lon", "longitude")
for need, name in ((c_alt, "altitudeRelative"), (c_air, "airSpeed"), (c_gs, "groundSpeed"), (c_lat, "gps.lat"), (c_lon, "gps.lon")):
    if need is None: sys.exit(f"У файлі немає стовпця {name}")
num = lambda c: pd.to_numeric(df[c], errors="coerce") if c else pd.Series([None]*len(df))
alt, air, gs = num(c_alt), num(c_air), num(c_gs)
roll, pitch, lat, lon = num(c_roll), num(c_pitch), num(c_lat), num(c_lon)
if c_ts:
    t = (pd.to_datetime(df[c_ts]) - pd.to_datetime(df[c_ts]).iloc[0]).dt.total_seconds().round(1)
else:
    t = pd.Series(range(len(df)), dtype=float)
n = len(df); last = n + 1                       # останній рядок даних в Excel

# межі етапів (для допоміжних стовпців графіків; у таблиці вони рахуються формулами)
near = (alt >= 0.98*H).to_numpy()
i_to = int(near.argmax()); i_ld = n - 1 - int(near[::-1].argmax())

ARIAL = "Arial"
f_norm = Font(name=ARIAL, size=10); f_bold = Font(name=ARIAL, size=10, bold=True)
f_in = Font(name=ARIAL, size=10, color="0000FF")
hdr_fill = PatternFill("solid", fgColor="D9E2F3")
thin = Side(style="thin", color="999999"); box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
wsd = wb.active; wsd.title = "Дані"
wst = wb.create_sheet("Таблиця 7.4")
wsg = wb.create_sheet("Графіки")

# ---------- Таблиця 7.4: параметри ----------
wst["A1"] = "Таблиця 7.4 – Результати аналізу телеметрії (контрольний політ, варіант 9)"; wst["A1"].font = Font(name=ARIAL, size=12, bold=True)
params = [("Задана висота H, м", H), ("Задана швидкість V, м/с", V), ("Цільова широта, °", TARGET[0]), ("Цільова довгота, °", TARGET[1])]
for i, (k, v) in enumerate(params, start=3):
    wst.cell(i, 1, k).font = f_norm; c = wst.cell(i, 2, v); c.font = f_in
wst["C3"] = "Вхідні параметри варіанта 9 (таблиця 7.6 методички); синій колір – значення, введені вручну"
wst["C3"].font = Font(name=ARIAL, size=9, italic=True)
# B3=H B4=V B5=lat B6=lon

D = "Дані"
rng = lambda L: f"{D}!${L}$2:${L}${last}"
rows = [
    ("Показник", "Значення", "Одиниця"),
    ("Час зльоту (до досягнення 98 % H)", f"=INDEX({rng('A')},MATCH(1,{rng('K')},0))", "с"),
    ("Початок посадки (останній момент на висоті ≈ H)", f"=SUMPRODUCT(MAX({rng('K')}*{rng('A')}))", "с"),
    ("Найбільша висота (відносна)", f"=MAX({rng('C')})", "м"),
    ("Найбільша наземна швидкість", f"=MAX({rng('E')})", "м/с"),
    ("Найбільша повітряна швидкість", f"=MAX({rng('D')})", "м/с"),
    ("Середня наземна швидкість на крейсері", f"=AVERAGEIFS({rng('E')},{rng('A')},\">\"&(B8+25),{rng('A')},\"<\"&(B9-20))", "м/с"),
    ("Відхилення середньої швидкості від заданої", "=ABS(B13-B4)/B4", "%"),
    ("Час польоту (arm → disarm)", f"=MAX({rng('A')})", "с"),
    ("Початкова широта", f"={D}!H2", "°"),
    ("Початкова довгота", f"={D}!I2", "°"),
    ("Кінцева широта", f"={D}!H{last}", "°"),
    ("Кінцева довгота", f"={D}!I{last}", "°"),
    ("Відстань кінцевої точки від цілі (наближено)", "=SQRT(((B18-B5)*111320)^2+((B19-B6)*111320*COS(RADIANS(B5)))^2)", "м"),
    ("Максимальний roll (за модулем)", f"=MAX(MAX({rng('F')}),-MIN({rng('F')}))", "град."),
    ("Максимальний pitch (за модулем)", f"=MAX(MAX({rng('G')}),-MIN({rng('G')}))", "град."),
    ("Режими польоту", "Auto (Takeoff → Waypoint → Land)", ""),
    ("Стан arm до та після місії", "disarmed → armed → disarmed", ""),
]
r0 = 7
for i, (a, b, c) in enumerate(rows):
    r = r0 + i
    wst.cell(r, 1, a); wst.cell(r, 2, b); wst.cell(r, 3, c)
    for cc in (1, 2, 3):
        cell = wst.cell(r, cc); cell.border = box; cell.font = f_bold if i == 0 else f_norm
        if i == 0: cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
# рядки: 7 шапка, 8 час зльоту, 9 початок посадки, 10 maxH, 11 maxGS, 12 maxAS, 13 mean, 14 відх., 15 час польоту,
# 16-17 поч. коорд., 18-19 кін. коорд., 20 відстань, 21 roll, 22 pitch, 23 режим, 24 arm
wst["B14"].number_format = "0.0%"; wst["C14"] = ""
for r in (8, 9, 10, 11, 12, 13, 15, 20, 21, 22): wst.cell(r, 2).number_format = "0.0"
for r in (16, 17, 18, 19): wst.cell(r, 2).number_format = "0.0000000"
wst["B13"].number_format = "0.00"; wst["B11"].number_format = "0.00"; wst["B12"].number_format = "0.00"
wst.column_dimensions["A"].width = 52; wst.column_dimensions["B"].width = 34; wst.column_dimensions["C"].width = 12
wst["A26"] = ("Примітки: усі числові показники обчислюються формулами за аркушем «Дані». Режим і стан arm у CSV відсутні "
              "(взято з індикаторів QGC). Відстань до цілі – наближена (1° ≈ 111 320 м).")
wst["A26"].font = Font(name=ARIAL, size=9, italic=True); wst["A26"].alignment = Alignment(wrap_text=True)
wst.merge_cells("A26:C27"); wst.row_dimensions[26].height = 30

# ---------- Дані ----------
heads = ["Час від arm, с", "Timestamp", "Висота (altitudeRelative), м", "Повітряна швидкість (airSpeed), м/с",
         "Наземна швидкість (groundSpeed), м/с", "roll, град.", "pitch, град.", "Широта, °", "Довгота, °",
         "Етап польоту", "Висота ≥ 98 % H", "Задана H, м", "Задана V, м/с",
         "Висота – зліт", "Висота – політ", "Висота – посадка",
         "Повітр. швидк. – зліт", "Повітр. швидк. – політ", "Повітр. швидк. – посадка",
         "Наземна швидк. – зліт", "Наземна швидк. – політ", "Наземна швидк. – посадка"]
for j, h in enumerate(heads, start=1):
    c = wsd.cell(1, j, h); c.font = f_bold; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = box
wsd.row_dimensions[1].height = 45
T = "'Таблиця 7.4'"
def seg(i):                                    # 0 = зліт, 1 = політ, 2 = посадка (з перекриттям на межах)
    s = []
    if i <= i_to: s.append(0)
    if i_to <= i <= i_ld: s.append(1)
    if i >= i_ld: s.append(2)
    return s
for i in range(n):
    r = i + 2
    vals = [t.iloc[i], str(df[c_ts].iloc[i]) if c_ts else "", alt.iloc[i], air.iloc[i], gs.iloc[i], roll.iloc[i], pitch.iloc[i], lat.iloc[i], lon.iloc[i]]
    for j, v in enumerate(vals, start=1):
        if pd.isna(v) if not isinstance(v, str) else False: v = None
        wsd.cell(r, j, v)
    wsd.cell(r, 10, f'=IF(A{r}<={T}!$B$8,"Зліт",IF(A{r}<={T}!$B$9,"Горизонтальний політ","Посадка"))')
    wsd.cell(r, 11, f"=IF(C{r}>=0.98*{T}!$B$3,1,0)")
    wsd.cell(r, 12, f"={T}!$B$3"); wsd.cell(r, 13, f"={T}!$B$4")
    for k, series in enumerate((alt, air, gs)):
        for ph in seg(i):
            v = series.iloc[i]
            wsd.cell(r, 14 + 3*k + ph, None if pd.isna(v) else float(v))
for row in wsd.iter_rows(min_row=2, max_row=last):
    for c in row: c.font = f_norm
for j in range(1, 23): wsd.column_dimensions[chr(64+j)].width = 16
wsd.column_dimensions["B"].width = 24; wsd.column_dimensions["J"].width = 22
wsd.freeze_panes = "A2"
wsd["N1"].comment = Comment("Стовпці N–V – допоміжні: значення розкладені за етапами польоту, щоб на графіках "
                            "етапи мали різні кольори. Межі етапів: час зльоту та початок посадки з «Таблиця 7.4» (B8, B9).", "lab7")

# ---------- Графіки ----------
def make_chart(title, ytitle, first_col, ref_col, ref_name, anchor, ymax):
    ch = ScatterChart(); ch.title = title; ch.style = 2; ch.height = 9.5; ch.width = 24
    ch.x_axis.title = "Час від arm, с"; ch.y_axis.title = ytitle
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.x_axis.scaling.min = 0; ch.x_axis.scaling.max = float(t.max()); ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = ymax
    ch.x_axis.majorGridlines = None
    xs = Reference(wsd, min_col=1, min_row=2, max_row=last)
    for k, (name, color) in enumerate((("Зліт", "1F77B4"), ("Горизонтальний політ", "2CA02C"), ("Посадка", "FF7F0E"))):
        s = Series(Reference(wsd, min_col=first_col + k, min_row=2, max_row=last), xs, title=name)
        s.marker.symbol = "none"; s.smooth = False
        s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = 28000
        ch.series.append(s)
    s = Series(Reference(wsd, min_col=ref_col, min_row=2, max_row=last), xs, title=ref_name)
    s.marker.symbol = "none"; s.smooth = False
    s.graphicalProperties.line.solidFill = "FF0000"; s.graphicalProperties.line.dashStyle = "dash"; s.graphicalProperties.line.width = 15000
    ch.series.append(s)
    ch.legend.position = "b"
    wsg.add_chart(ch, anchor)
make_chart("Графік 1 – Відносна висота від часу", "Висота, м", 14, 12, "Задана H", "A1", 40)
make_chart("Графік 2 – Повітряна швидкість від часу", "Швидкість, м/с", 17, 13, "Задана V", "A22", 5)
make_chart("Графік 3 – Наземна швидкість від часу", "Швидкість, м/с", 20, 13, "Задана V", "A43", 5)

os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
wb.save(out)
print("Збережено:", os.path.abspath(out), f"(рядків даних: {n}; зліт до {t.iloc[i_to]:.0f} с, посадка з {t.iloc[i_ld]:.0f} с)")
